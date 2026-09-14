import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from django.test import TestCase
from ge.views_utils import (
    create_excel_output,
    get_as_of_date,
    get_last_col,
    get_last_row,
    sum_col,
)


class ReportManipulationTestCase(TestCase):
    # sample file, structured like a Unit report but without sums
    sample_report = "sample_files/sample_report.xlsx"
    wb = load_workbook(sample_report)
    gifts_ws = wb["Gifts"]

    def test_last_row(self):
        # Last row in sample report is 6 on non-summed columns, like Unit (A)
        self.assertEqual(get_last_row(self.gifts_ws, "A"), 6)

    def test_last_col(self):
        # Last column in sample report is S (19), and data is in row 5
        self.assertEqual(get_last_col(self.gifts_ws, 5), 19)

    def test_sum_col(self):
        # add a sum in column L (YTD Appropriation)
        sum_col(self.gifts_ws, "L")
        self.assertEqual(self.gifts_ws["L7"].value, "=SUM(L5:L6)")

    def test_as_of_date_current_year(self):
        # In November, as-of date should be 9/30 of current year
        self.assertEqual(get_as_of_date(datetime(2023, 11, 1)), "as of 9/30/23")

    def test_as_of_date_previous_year(self):
        # In February, as-of date should be 12/31 of previous year
        self.assertEqual(get_as_of_date(datetime(2024, 2, 1)), "as of 12/31/23")


class ExcelOutputTestCase(TestCase):
    @classmethod
    def setUpTestData(cls) -> None:
        # Create DataFrames from real saved data, combining local and
        # QDB data as of 202607.
        aul_endowments_df = pd.read_csv("ge/fixtures/aul_endowments.csv")
        aul_gifts_df = pd.read_csv("ge/fixtures/aul_gifts.csv")
        ul_endowments_df = pd.read_csv("ge/fixtures/ul_endowments.csv")
        ul_gifts_df = pd.read_csv("ge/fixtures/ul_gifts.csv")
        unit_endowments_df = pd.read_csv("ge/fixtures/unit_endowments.csv")
        unit_gifts_df = pd.read_csv("ge/fixtures/unit_gifts.csv")

        # Master reports use only one set of data.
        master_df = pd.read_csv("ge/fixtures/master.csv")
        cls.master_data = (master_df, pd.DataFrame())

        # The others use 2 sets.
        cls.aul_data = (aul_endowments_df, aul_gifts_df)
        cls.ul_data = (ul_endowments_df, ul_gifts_df)
        cls.unit_data = (unit_endowments_df, unit_gifts_df)

    def test_master_report_worksheets(self):
        result = create_excel_output("master", self.master_data)
        # Only one worksheet in Master report
        self.assertEqual(len(result.sheetnames), 1)

    def test_master_report_cols(self):
        result = create_excel_output("master", self.master_data)
        # Last column is "LBS Notes" in column U
        self.assertEqual(result["G&E"]["U2"].value, "LBS Notes")

    def test_master_report_rows(self):
        result = create_excel_output("master", self.master_data)
        # 13 rows in sample data. Data starts on row 5, so we should have data
        # in rows 5-17 and not in 18.
        # Master report isn't sorted, so just check if data exists
        self.assertNotEqual(result["G&E"]["A17"].value, None)
        self.assertEqual(result["G&E"]["A18"].value, None)

    def test_unit_report_worksheets(self):
        result = create_excel_output("arts", self.unit_data)
        # two worksheets in Arts report
        self.assertEqual(len(result.sheetnames), 2)

    def test_unit_report_cols(self):
        result = create_excel_output("arts", self.unit_data)
        # Arts sample data has fund restriction for Endowments, but not Gifts
        # So "Fund Restriction" should be in column S for Endowments,
        # and Gifts should have "LBS Notes" in column S
        self.assertEqual(result["Endowments"]["S2"].value, "Fund Restriction")
        self.assertEqual(result["Gifts"]["S2"].value, "LBS Notes")

    def test_unit_report_rows(self):
        result = create_excel_output("arts", self.unit_data)
        # Arts data has 1 gift and 2 endowments, starting on row 5
        self.assertEqual(result["Gifts"]["A5"].value, "Arts")
        self.assertEqual(result["Gifts"]["A6"].value, None)
        self.assertEqual(result["Endowments"]["A5"].value, "Arts")
        self.assertEqual(result["Endowments"]["A7"].value, None)

    def test_unit_report_totals(self):
        result = create_excel_output("arts", self.unit_data)
        # Gifts should have totals in cols L, M, N, O. Endowments in L, M, N, O, Q
        # SUM is calculated by Excel, so just check the formulas
        self.assertEqual(result["Gifts"]["L6"].value, "=SUM(L5:L5)")
        self.assertEqual(result["Gifts"]["M6"].value, "=SUM(M5:M5)")
        self.assertEqual(result["Gifts"]["N6"].value, "=SUM(N5:N5)")
        self.assertEqual(result["Gifts"]["O6"].value, "=SUM(O5:O5)")
        self.assertEqual(result["Endowments"]["L7"].value, "=SUM(L5:L6)")
        self.assertEqual(result["Endowments"]["M7"].value, "=SUM(M5:M6)")
        self.assertEqual(result["Endowments"]["N7"].value, "=SUM(N5:N6)")
        self.assertEqual(result["Endowments"]["O7"].value, "=SUM(O5:O6)")
        self.assertEqual(result["Endowments"]["Q7"].value, "=SUM(Q5:Q6)")

    def test_ul_report_worksheets(self):
        result = create_excel_output("ul", self.ul_data)
        # two worksheets in UL report
        self.assertEqual(len(result.sheetnames), 2)

    def test_ul_report_cols(self):
        result = create_excel_output("ul", self.ul_data)
        # No fund restrictions, so "LBS Notes" should be in column T for Endowments, S for Gifts
        self.assertEqual(result["Endowments"]["T2"].value, "LBS Notes")
        self.assertEqual(result["Gifts"]["S2"].value, "LBS Notes")

    def test_ul_report_rows(self):
        result = create_excel_output("ul", self.ul_data)
        # UL data has 1 gift and 1 endowment, starting on row 5
        self.assertEqual(result["Gifts"]["A5"].value, "UL")
        self.assertEqual(result["Gifts"]["A6"].value, None)
        self.assertEqual(result["Endowments"]["A5"].value, "UL")
        self.assertEqual(result["Endowments"]["A6"].value, None)

    def test_ul_report_totals(self):
        result = create_excel_output("ul", self.ul_data)
        # Gifts should have totals in cols L, M, N, O. Endowments in L, M, N, O, Q.
        # SUM is calculated by Excel, so just check the formulas
        self.assertEqual(result["Gifts"]["L6"].value, "=SUM(L5:L5)")
        self.assertEqual(result["Gifts"]["M6"].value, "=SUM(M5:M5)")
        self.assertEqual(result["Gifts"]["N6"].value, "=SUM(N5:N5)")
        self.assertEqual(result["Gifts"]["O6"].value, "=SUM(O5:O5)")
        self.assertEqual(result["Endowments"]["L6"].value, "=SUM(L5:L5)")
        self.assertEqual(result["Endowments"]["M6"].value, "=SUM(M5:M5)")
        self.assertEqual(result["Endowments"]["N6"].value, "=SUM(N5:N5)")
        self.assertEqual(result["Endowments"]["O6"].value, "=SUM(O5:O5)")
        self.assertEqual(result["Endowments"]["Q6"].value, "=SUM(Q5:Q5)")

    def test_aul_report_worksheets(self):
        result = create_excel_output("aul_benedetti", self.aul_data)
        # two worksheets in AUL report
        self.assertEqual(len(result.sheetnames), 2)

    def test_aul_report_cols(self):
        result = create_excel_output("aul_benedetti", self.aul_data)
        # AUL sample data has fund restriction for Gifts, but not Endowments.
        # So "Fund Restriction" should be in column R for Gifts,
        # and Endowments should have "LBS Notes" in column T
        self.assertEqual(result["Gifts"]["R2"].value, "Fund Restriction")
        self.assertEqual(result["Endowments"]["T2"].value, "LBS Notes")

    def test_aul_report_rows(self):
        result = create_excel_output("aul_benedetti", self.aul_data)
        # 1 gift and 1 endowment, starting on row 5
        # Fuzzy match is used for AUL reports, against two columns.
        # For Gifts, value is in unit name (column A).
        self.assertIn("Benedetti", result["Gifts"]["A5"].value)
        # For Endowments, value is in home / unit dept but not unit
        # (which is populated here, but not relevant).
        self.assertIn("Benedetti", result["Endowments"]["B5"].value)

    def test_aul_report_totals(self):
        result = create_excel_output("aul_benedetti", self.aul_data)
        # Gifts should have totals in cols L, M, N, O. Endowments in L, M, N, O, Q
        # SUM is calculated by Excel, so just check the formulas
        self.assertEqual(result["Gifts"]["L6"].value, "=SUM(L5:L5)")
        self.assertEqual(result["Gifts"]["M6"].value, "=SUM(M5:M5)")
        self.assertEqual(result["Gifts"]["N6"].value, "=SUM(N5:N5)")
        self.assertEqual(result["Gifts"]["O6"].value, "=SUM(O5:O5)")
        self.assertEqual(result["Endowments"]["L6"].value, "=SUM(L5:L5)")
        self.assertEqual(result["Endowments"]["M6"].value, "=SUM(M5:M5)")
        self.assertEqual(result["Endowments"]["N6"].value, "=SUM(N5:N5)")
        self.assertEqual(result["Endowments"]["O6"].value, "=SUM(O5:O5)")
        self.assertEqual(result["Endowments"]["Q6"].value, "=SUM(Q5:Q5)")
